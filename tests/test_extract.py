"""Tests for tidy_one.extract: scanning, sensitivity, extraction and chunking. All fixtures are generated."""

from __future__ import annotations

import datetime as dt
import importlib
import os
import random
import stat
import struct
import unicodedata
import zipfile
import zlib
from pathlib import Path
from xml.sax.saxutils import escape

import pytest

from tidy_one.embed.fake import FakeEncoder
from tidy_one.extract import chunk_blocks, content_sensitivity, extract_file, is_secret_path, scan, split_text
from tidy_one.extract.extract import TextBlock, decode_bytes, hwp_records_text, rtf_to_text
from tidy_one.extract.sensitive import find_card_numbers, find_passports, find_rrns, luhn_ok, valid_rrn
from tidy_one.model import SENS_ID, SENS_NONE, SENS_SECRET

BS = chr(92)  # backslash

scan_mod = importlib.import_module("tidy_one.extract.scan")  # the package re-exports a function named `scan`
extract_mod = importlib.import_module("tidy_one.extract.extract")


# ---------------------------------------------------------------- fixture builders


def touch(p: Path, data: bytes | str = "x") -> Path:
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(data.encode() if isinstance(data, str) else data)
    return p


def make_docx(path: Path, paras: list[str], title: str) -> None:
    ct = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        '<Override PartName="/docProps/core.xml" '
        'ContentType="application/vnd.openxmlformats-package.core-properties+xml"/></Types>'
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
        'officeDocument" Target="word/document.xml"/>'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/'
        'core-properties" Target="docProps/core.xml"/></Relationships>'
    )
    body = "".join(f"<w:p><w:r><w:t>{escape(p)}</w:t></w:r></w:p>" for p in paras)
    doc = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{body}</w:body></w:document>"
    )
    core = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
        f'xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>{escape(title)}</dc:title></cp:coreProperties>'
    )
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", ct)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", doc)
        z.writestr("docProps/core.xml", core)


def make_hwpx(path: Path, sections: list[str], title: str) -> None:
    hpf = (
        '<?xml version="1.0" encoding="UTF-8"?><opf:package xmlns:opf="http://www.idpf.org/2007/opf/">'
        f"<opf:metadata><opf:title>{escape(title)}</opf:title></opf:metadata></opf:package>"
    )
    head = (
        '<?xml version="1.0" encoding="UTF-8"?><hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" '
        'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">'
    )
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("mimetype", "application/hwp+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("Contents/content.hpf", hpf)
        for i, body in enumerate(sections):
            z.writestr(f"Contents/section{i}.xml", head + body + "</hs:sec>")
        z.writestr("Preview/PrvText.txt", "미리보기")


def make_pdf(path: Path, pages: list[str], title: str | None = None) -> None:
    """Minimal PDF with one Helvetica text line per page (ASCII only)."""
    n = len(pages)
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{' '.join(f'{4 + 2 * i} 0 R' for i in range(n))}] /Count {n} >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for i, text in enumerate(pages):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objs.append(
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>"
        )
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    if title:
        objs.append(f"<< /Title ({title}) >>")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for k, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{k} 0 obj\n{o}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{off:010d} 00000 n \n".encode() for off in offsets)
    info = f" /Info {len(objs)} 0 R" if title else ""
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R{info} >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))


_END, _FREE, _FATSECT, _NOSTREAM = 0xFFFFFFFE, 0xFFFFFFFF, 0xFFFFFFFD, 0xFFFFFFFF


def make_cfb(tree: dict[str, bytes | dict[str, bytes]]) -> bytes:
    """Minimal OLE2 compound file (v3, 512-byte sectors, one FAT sector, no mini stream)."""
    entries: list[dict] = [{"name": "Root Entry", "type": 5, "kids": []}]

    def add(parent: int, name: str, val: bytes | dict[str, bytes]) -> None:
        idx = len(entries)
        if isinstance(val, dict):
            entries.append({"name": name, "type": 1, "kids": []})
            for n, v in val.items():
                add(idx, n, v)
        else:
            entries.append({"name": name, "type": 2, "kids": [], "data": val.ljust(4096, b"\0")})
        entries[parent]["kids"].append(idx)

    for n, v in tree.items():
        add(0, n, v)
    ndir = (len(entries) * 128 + 511) // 512
    fat = [_FREE] * 128
    fat[0] = _FATSECT
    for i in range(1, ndir + 1):
        fat[i] = i + 1 if i < ndir else _END
    nxt, body = ndir + 1, b""
    for e in entries:
        e.update(left=_NOSTREAM, right=_NOSTREAM, child=_NOSTREAM)
        if e["type"] == 2:
            k = len(e["data"]) // 512
            e["start"] = nxt
            for j in range(k):
                fat[nxt + j] = nxt + j + 1 if j < k - 1 else _END
            body += e["data"]
            nxt += k
    assert nxt <= 128
    for e in entries:
        kids = sorted(e["kids"], key=lambda i: (len(entries[i]["name"]), entries[i]["name"].upper()))
        if kids:
            e["child"] = kids[0]
            for a, b in zip(kids, kids[1:]):
                entries[a]["right"] = b
    dirs = b""
    for e in entries:
        name = e["name"].encode("utf-16-le") + b"\0\0"
        dirs += name.ljust(64, b"\0")
        dirs += struct.pack("<HBBIII", len(name), e["type"], 1, e["left"], e["right"], e["child"])
        dirs += b"\0" * 36
        start = e.get("start", _END if e["type"] == 5 else 0)
        dirs += struct.pack("<IQ", start, len(e.get("data", b"")))
    dirs = dirs.ljust(ndir * 512, b"\0")
    header = struct.pack(
        "<8s16sHHHHH6sIIIIIIIII",
        b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", b"\0" * 16, 0x3E, 3, 0xFFFE, 9, 6, b"\0" * 6,
        0, 1, 1, 0, 4096, _END, 0, _END, 0,
    )
    header += struct.pack("<109I", 0, *([_FREE] * 108))
    return header + struct.pack("<128I", *fat) + dirs + body


def hwp_record(tag: int, payload: bytes) -> bytes:
    return struct.pack("<I", tag | (len(payload) << 20)) + payload


def hwp_para(text: str) -> bytes:
    """PARA_TEXT payload: text, an extended control (8 WCHARs), a tab (8 WCHARs) and the paragraph end."""
    ext_ctrl = struct.pack("<H", 11) + b"dces" + b"\0" * 8 + struct.pack("<H", 11)
    tab = struct.pack("<H", 9) + b"\0" * 12 + struct.pack("<H", 9)
    return text.encode("utf-16-le") + ext_ctrl + tab + "끝".encode("utf-16-le") + struct.pack("<H", 13)


def make_hwp(path: Path, paras: list[str], *, compressed: bool = True, distribution: bool = False) -> None:
    props = (1 if compressed else 0) | (4 if distribution else 0)
    header = b"HWP Document File".ljust(32, b"\0") + struct.pack("<II", 0x05000300, props)
    body = b"".join(hwp_record(66, b"\0" * 22) + hwp_record(67, hwp_para(p)) for p in paras)
    if compressed:
        c = zlib.compressobj(wbits=-15)
        body = c.compress(body) + c.flush()
    path.write_bytes(make_cfb({
        "FileHeader": header.ljust(256, b"\0"),
        "BodyText": {"Section0": body},
        "PrvText": "미리보기 텍스트입니다".encode("utf-16-le"),
    }))


def rrn_with_check(front6: str, back6: str) -> str:
    d = front6 + back6
    w = (2, 3, 4, 5, 6, 7, 8, 9, 2, 3, 4, 5)
    check = (11 - sum(int(a) * b for a, b in zip(d, w)) % 11) % 10
    return f"{front6}-{back6}{check}"


# ---------------------------------------------------------------- scan


def _rel(root: Path, infos) -> set[str]:
    return {os.path.relpath(i.path, root) for i in infos}


def test_scan_skip_rules(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    home = tmp_path.resolve() / "home"
    monkeypatch.setenv("HOME", str(home))
    root = home
    keep = [
        "Documents/보고서.txt",
        "Documents/notes/build/plan.md",
        "Desktop/a.pdf",
        "Documents/proj/package.json",
    ]
    skip = [
        ".hidden.txt",
        ".config/app.json",
        "Library/Preferences/x.plist",
        "Documents/node_modules/pkg/index.js",
        "Documents/__pycache__/m.pyc",
        "Documents/myrepo/.git/HEAD",
        "Documents/myrepo/README.md",
        "Documents/proj/dist/bundle.js",
        "Documents/env/pyvenv.cfg",
        "Documents/env/lib/site.py",
        "Applications/Foo.app/Contents/Info.plist",
        "Pictures/Photos Library.photoslibrary/database/x.db",
        "Documents/~$report.docx",
        "Documents/passwords.txt",
        "Documents/id_rsa",
        "Documents/secrets/plan.txt",
        "Documents/old.bak",
        "Documents/skipme/a.txt",
        "Documents/drafts/old/a.txt",
        "Documents/DerivedData/x.o",
        "Documents/work/font-cache/CACHEDIR.TAG",
        "Documents/work/font-cache/abc-le64.cache-9",
    ]
    for p in keep + skip:
        touch(root / p)
    touch(root / "Documents/drafts/new.txt")
    keep.append("Documents/drafts/new.txt")
    os.symlink(root / "Documents/보고서.txt", root / "Documents/link.txt")
    os.symlink(root / "Documents", root / "docs-link")
    with open(root / "Documents/huge.bin", "wb") as f:
        f.truncate(201 * 1024 * 1024)  # sparse; over the 200MB cap

    got = _rel(root, scan([str(root)], ["*.bak", "skipme", "Documents/drafts/old/**"]))
    assert got == set(keep)


def test_scan_explicit_root_repo_and_skip_repos_flag(tmp_path: Path) -> None:
    tmp_path = tmp_path.resolve()
    repo = tmp_path / "code" / "repo"
    touch(repo / ".git" / "HEAD")
    touch(repo / "README.md")
    touch(repo / "sub" / "x.txt")
    assert _rel(tmp_path, scan([str(tmp_path)])) == set()
    assert _rel(tmp_path, scan([str(repo)])) == {"code/repo/README.md", "code/repo/sub/x.txt"}
    assert _rel(tmp_path, scan([str(tmp_path)], skip_repos=False)) == {"code/repo/README.md", "code/repo/sub/x.txt"}


def test_scan_dataless_and_hidden_flag(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    tmp_path = tmp_path.resolve()
    touch(tmp_path / "local.txt")
    touch(tmp_path / "cloud.txt")
    touch(tmp_path / "clouddir" / "inside.txt")
    touch(tmp_path / "flagged.txt")
    if hasattr(os, "chflags"):
        os.chflags(tmp_path / "flagged.txt", stat.UF_HIDDEN)
    else:
        (tmp_path / "flagged.txt").unlink()
    real = scan_mod._st_flags

    def fake_flags(st: os.stat_result, path: str) -> int:
        name = os.path.basename(path)
        return real(st, path) | (scan_mod.SF_DATALESS if name in ("cloud.txt", "clouddir") else 0)

    monkeypatch.setattr(scan_mod, "_st_flags", fake_flags)
    assert scan_mod.SF_DATALESS == 0x40000000
    assert _rel(tmp_path, scan([str(tmp_path)])) == {"local.txt"}


def test_scan_overlapping_roots_and_file_root(tmp_path: Path) -> None:
    tmp_path = tmp_path.resolve()
    touch(tmp_path / "docs" / "a.md", "hello")
    walked = list(scan([str(tmp_path / "docs"), str(tmp_path)]))
    assert [os.path.relpath(i.path, tmp_path) for i in walked] == ["docs/a.md"]
    info = walked[0]
    assert info.size == 5 and info.mtime > 0 and info.ctime > 0
    single = list(scan([str(tmp_path / "docs" / "a.md"), str(tmp_path / "missing")]))
    assert [i.path for i in single] == [str(tmp_path / "docs" / "a.md")]


def test_scan_order_is_deterministic(tmp_path: Path) -> None:
    tmp_path = tmp_path.resolve()
    for p in ["b/2.txt", "a/1.txt", "c.txt", "a/b/3.txt"]:
        touch(tmp_path / p)
    assert [os.path.relpath(i.path, tmp_path) for i in scan([str(tmp_path)])] == [
        "c.txt", "a/1.txt", "a/b/3.txt", "b/2.txt",
    ]


# ---------------------------------------------------------------- sensitivity


@pytest.mark.parametrize(
    "path,secret",
    [
        ("~/.env", True),
        ("/x/.env.local", True),
        ("/x/server.pem", True),
        ("/x/tls.key", True),
        ("/x/id_rsa", True),
        ("/x/id_ed25519.pub", True),
        ("/x/vault.kdbx", True),
        ("/x/cert.p12", True),
        ("/x/login.keychain-db", True),
        ("/x/My Passwords.xlsx", True),
        ("/x/aws_credentials.csv", True),
        ("/x/slack-token.txt", True),
        ("/x/client_secret_123.json", True),
        ("/x/은행 비밀번호.txt", True),
        ("/x/" + unicodedata.normalize("NFD", "암호 목록") + ".txt", True),
        ("/x/" + unicodedata.normalize("NFD", "비밀번호") + ".txt", True),
        ("/home/u/.ssh/config", True),
        ("/x/보고서.pdf", False),
        ("/x/meeting-notes.md", False),
        ("/x/keyboard.txt", False),
        ("/x/work/artcraft-host-key", True),
        ("/x/ssh_host_ed25519_key", True),
        ("/x/deploy_key", True),
        ("/x/api-keys.txt", True),
        ("/x/hotkeys.md", False),
        ("/x/keynote-tips.md", False),
        ("/x/host-key.pub", False),
    ],
)
def test_is_secret_path(path: str, secret: bool) -> None:
    assert is_secret_path(path) is secret


def test_rrn_checksum() -> None:
    good = rrn_with_check("900101", "123456")
    bad = good[:-1] + str((int(good[-1]) + 1) % 10)
    assert valid_rrn(good) and not valid_rrn(bad)
    assert find_rrns(f"주민등록번호: {good} 입니다") == [good]
    assert find_rrns(f"주민등록번호: {bad}") == []
    assert content_sensitivity(f"성명 홍길동 {good}") == SENS_ID
    assert content_sensitivity(f"번호 {bad}") == SENS_NONE
    assert not valid_rrn("901301-1234567")  # month 13
    assert not valid_rrn("900101-5234567")  # gender digit outside 1-4
    assert find_rrns("주문번호 12900101-1234567") == []  # embedded in a longer number
    # issued from October 2020: trailing digits are random, no checksum
    assert valid_rrn("210315-3999999") and valid_rrn("201001-4000001")
    post = rrn_with_check("200915", "312345")
    assert not valid_rrn(post[:-1] + str((int(post[-1]) + 1) % 10))


def test_card_and_passport() -> None:
    assert luhn_ok("4111111111111111") and not luhn_ok("4111111111111112")
    assert find_card_numbers("카드 4111 1111 1111 1111 결제") == ["4111111111111111"]
    assert find_card_numbers("카드 4111-1111-1111-1112") == []
    assert find_card_numbers("주문번호 1234567890123456") == []  # no card brand prefix
    assert find_card_numbers("아멕스 3782-822463-10005") == ["378282246310005"]
    assert find_card_numbers('"digest":"sha256:ab4111111111111111c9"') == []  # inside a hex digest
    assert find_card_numbers("카드번호4111111111111111입니다") == ["4111111111111111"]
    assert find_passports("여권번호: M12345678") == ["M12345678"]
    assert find_passports("Passport No. M123A4567") == ["M123A4567"]
    assert find_passports("모델명 M12345678 제품") == []
    assert content_sensitivity("여권 번호 M12345678") == SENS_ID
    assert content_sensitivity("결제 카드 5555 5555 5555 4444") == SENS_ID


def test_secret_content_patterns() -> None:
    samples = [
        "-----BEGIN OPENSSH " + "PRIVATE KEY-----\nabc\n-----END OPENSSH PRIVATE KEY-----",
        "-----BEGIN RSA " + "PRIVATE KEY-----",
        "key=" + "AKIA" + "IOSFODNN7EXAMPLE",
        "aws_secret_access_key = " + "wJalrXUtnFEMI/K7MDENG/bPxRfiCYEXAMPLEKEY",
        "ANTHROPIC_API_KEY=" + "sk-ant-" + "api03-" + "A" * 40,
        "OPENAI=" + "sk-" + "proj-" + "B" * 40,
        "gh" + "p_" + "C" * 36,
        "github_" + "pat_" + "D" * 30,
        "xo" + "xb-" + "1234567890-abcdefghij",
        "AI" + "za" + "E" * 35,
        '{"type": "service_account", "project_id": "x", "private_key": "..."}',
        "ntn" + "_" + "F" * 46,
    ]
    for s in samples:
        assert content_sensitivity(s) & SENS_SECRET, s
    assert content_sensitivity("sk-learn은 파이썬 라이브러리, AKIA는 접두어") == SENS_NONE


# ---------------------------------------------------------------- extraction


def test_text_encodings(tmp_path: Path) -> None:
    ko = "안녕하세요. 한국어 문서입니다.\n둘째 줄, 숫자 2026."
    cases = {
        "cp949.txt": (ko.encode("cp949"), "cp949"),
        "utf8.txt": (ko.encode("utf-8"), "utf-8"),
        "bom8.txt": (b"\xef\xbb\xbf" + ko.encode("utf-8"), "utf-8-sig"),
        "u16.txt": (ko.encode("utf-16"), "utf-16"),
        "u16le.txt": (ko.encode("utf-16-le"), "utf-16-le"),
        "u16be.txt": (ko.encode("utf-16-be"), "utf-16-be"),
        "euckr.txt": (ko.encode("euc-kr"), "cp949"),
    }
    for name, (data, enc) in cases.items():
        ex = extract_file(str(touch(tmp_path / name, data)))
        assert ex.text == ko, name
        assert ex.extra["encoding"] == enc, name
        assert ex.modality == "text" and ex.mime == "text/plain"
        assert ex.title == name.rsplit(".", 1)[0]


def test_decode_partial_tail_is_tolerated() -> None:
    data = "가나다".encode("utf-8")[:-1]
    text, enc = decode_bytes(data, partial=True)
    assert enc == "utf-8" and text == "가나"


def test_markdown_title_and_csv_json(tmp_path: Path) -> None:
    md = touch(tmp_path / "note.md", "---\ntitle: 프로젝트 계획\ntags: [a]\n---\n\n# 다른 제목\n\n본문입니다.\n")
    ex = extract_file(str(md))
    assert ex.title == "프로젝트 계획" and "본문입니다." in ex.text and ex.mime == "text/markdown"
    md2 = touch(tmp_path / "n2.md", "intro\n\n# **회의** 요약\n\n내용")
    assert extract_file(str(md2)).title == "회의 요약"
    csv = touch(tmp_path / "list.csv", "이름,도시\n홍길동,서울\n김철수,부산\n".encode("cp949"))
    ex = extract_file(str(csv))
    assert ex.text.splitlines() == ["이름,도시", "홍길동,서울", "김철수,부산"] and ex.mime == "text/csv"
    js = touch(tmp_path / "d.json", '{"제목": "값"}')
    assert '"제목"' in extract_file(str(js)).text


def test_html_and_rtf(tmp_path: Path) -> None:
    html = touch(tmp_path / "p.html", "<html><head><title> 웹 페이지 </title><style>p{}</style></head>"
                 "<body><p>첫 문단</p><script>var x=1;</script><p>둘째 &amp; 문단</p></body></html>")
    ex = extract_file(str(html))
    assert ex.title == "웹 페이지" and "첫 문단" in ex.text and "둘째 & 문단" in ex.text and "var x" not in ex.text

    def u(s: str) -> str:  # one RTF unicode escape (backslash, 'u', decimal code point) per character
        return "".join(f"{BS}u{ord(c)} " for c in s)

    rtf = (
        r"{\rtf1\ansi\ansicpg949\deff0{\fonttbl{\f0 AppleSDGothicNeo;}}{\colortbl;\red0\green0\blue0;}"
        r"{\info{\title \uc0" + u("회의록") + r"}{\author X}}"
        r"{\*\expandedcolortbl;;}\f0 \uc0" + u("안녕하세요") + r". Hello\par "
        r"\'c7\'d1\'b1\'db \uc1" + u("글") + r"?\par}"
    )
    text, title = rtf_to_text(rtf)
    assert title == "회의록"
    assert text.splitlines() == ["안녕하세요. Hello", "한글 글"]
    ex = extract_file(str(touch(tmp_path / "r.rtf", rtf)))
    assert ex.title == "회의록" and "안녕하세요" in ex.text


def test_pdf_pages_and_needs_ocr(tmp_path: Path) -> None:
    from pypdf import PdfWriter

    pdf = tmp_path / "doc.pdf"
    make_pdf(pdf, ["Hello page one with enough words to count as text", "Second page about budgets and plans"],
             title="Quarterly Report")
    ex = extract_file(str(pdf))
    assert [b.page for b in ex.text_blocks] == [1, 2]
    assert "page one" in ex.text_blocks[0].text and "budgets" in ex.text_blocks[1].text
    assert ex.title == "Quarterly Report" and ex.extra["pages"] == 2 and "needs_ocr" not in ex.extra
    w = PdfWriter()
    w.add_blank_page(612, 792)
    w.add_blank_page(612, 792)
    scanned = tmp_path / "scan.pdf"
    with open(scanned, "wb") as f:
        w.write(f)
    ex = extract_file(str(scanned))
    assert ex.text_blocks == [] and ex.extra["needs_ocr"] is True and ex.title == "scan"


def test_office_formats(tmp_path: Path) -> None:
    docx = tmp_path / "회의록.docx"
    make_docx(docx, ["첫 문단입니다.", "둘째 문단 & 기호"], "주간 회의")
    ex = extract_file(str(docx))
    assert ex.title == "주간 회의" and "첫 문단입니다." in ex.text and "둘째 문단 & 기호" in ex.text

    from pptx import Presentation

    prs = Presentation()
    prs.core_properties.title = "발표 자료"
    for head, body in [("첫 슬라이드", "본문 하나"), ("둘째 슬라이드", "본문 둘")]:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = head
        s.placeholders[1].text = body
    pptx = tmp_path / "deck.pptx"
    prs.save(str(pptx))
    ex = extract_file(str(pptx))
    assert ex.title == "발표 자료" and [b.page for b in ex.text_blocks] == [1, 2]
    assert "본문 둘" in ex.text_blocks[1].text and "<!--" not in ex.text

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "매출"
    ws.append(["월", "금액"])
    ws.append(["10월", 1200])
    wb.create_sheet("비용").append(["서버", 30])
    xlsx = tmp_path / "book.xlsx"
    wb.save(xlsx)
    ex = extract_file(str(xlsx))
    assert ex.extra["sheets"] == ["매출", "비용"] and "10월" in ex.text and "서버" in ex.text
    assert ex.title == "book"


def test_hwpx(tmp_path: Path) -> None:
    p = tmp_path / "공문.hwpx"
    make_hwpx(p, [
        "<hp:p><hp:run><hp:t>첫 번째 문단입니다.</hp:t></hp:run></hp:p>"
        "<hp:p><hp:run><hp:t>앞<hp:tab/>뒤</hp:t></hp:run></hp:p>",
        "<hp:p><hp:run><hp:tbl><hp:tr><hp:tc><hp:subList><hp:p><hp:run><hp:t>셀 내용</hp:t></hp:run></hp:p>"
        "</hp:subList></hp:tc></hp:tr></hp:tbl></hp:run></hp:p>",
    ], "보도자료")
    ex = extract_file(str(p))
    assert ex.title == "보도자료" and ex.mime == "application/hwp+zip"
    assert ex.text.split("\n") == ["첫 번째 문단입니다.", "앞\t뒤", "", "셀 내용"]
    assert "preview_only" not in ex.extra


def test_hwp_binary(tmp_path: Path) -> None:
    p = tmp_path / "계획서.hwp"
    make_hwp(p, ["한글 문서 첫 문단", "둘째 문단"])
    ex = extract_file(str(p))
    assert "error" not in ex.extra, ex.extra
    assert ex.text.split("\n") == ["한글 문서 첫 문단\t끝", "둘째 문단\t끝"]
    assert ex.title == "계획서" and "preview_only" not in ex.extra

    plain = tmp_path / "plain.hwp"
    make_hwp(plain, ["압축 없는 문서"], compressed=False)
    assert extract_file(str(plain)).text.startswith("압축 없는 문서")

    dist = tmp_path / "배포용.hwp"
    make_hwp(dist, ["보이지 않는 본문"], distribution=True)
    ex = extract_file(str(dist))
    assert ex.text == "미리보기 텍스트입니다" and ex.extra["preview_only"] and ex.extra["hwp_protected"]


def test_hwp_records_long_header() -> None:
    payload = "긴 문단".encode("utf-16-le") * 700
    rec = struct.pack("<I", 67 | (0xFFF << 20)) + struct.pack("<I", len(payload)) + payload
    assert hwp_records_text(rec) == "긴 문단" * 700


def test_media_and_unknown_are_metadata_only(tmp_path: Path) -> None:
    for name, modality in [("IMG_0001.HEIC", "image"), ("photo.jpg", "image"), ("clip.mov", "video"),
                           ("memo.m4a", "audio"), ("archive.zip", "text"), ("noext", "text")]:
        ex = extract_file(str(touch(tmp_path / name, b"\x00\x01binary")))
        assert ex.modality == modality and ex.text_blocks == [] and ex.sensitivity == SENS_NONE, name
        assert ex.title == name.rsplit(".", 1)[0]


def test_secret_path_never_opened(tmp_path: Path) -> None:
    p = touch(tmp_path / "my passwords.txt", "hunter2")
    p.chmod(0)
    try:
        ex = extract_file(str(p))
    finally:
        p.chmod(0o600)
    assert ex.sensitivity == SENS_SECRET and ex.text_blocks == [] and ex.extra == {"skipped": "secret_path"}


def test_secret_and_id_content(tmp_path: Path) -> None:
    key = touch(tmp_path / "notes.txt", "서버 접속\n-----BEGIN OPENSSH " + "PRIVATE KEY-----\nxyz\n")
    ex = extract_file(str(key))
    assert ex.sensitivity & SENS_SECRET and ex.text_blocks == [] and ex.extra["skipped"] == "secret_content"
    rrn = rrn_with_check("850505", "234567")
    idoc = touch(tmp_path / "계약서.txt", f"계약자 주민등록번호 {rrn}\n계약 내용입니다.")
    ex = extract_file(str(idoc))
    assert ex.sensitivity == SENS_ID and "계약 내용" in ex.text


def test_text_cap_and_broken_files(tmp_path: Path) -> None:
    big = touch(tmp_path / "big.txt", "가나다라 " * 1000)
    ex = extract_file(str(big), max_chars=1000)
    assert len(ex.text) == 1000 and ex.extra["truncated"] is True
    broken = touch(tmp_path / "broken.docx", b"not a zip")
    ex = extract_file(str(broken))
    assert ex.text_blocks == [] and "error" in ex.extra and ex.title == "broken"
    bad_pdf = touch(tmp_path / "bad.pdf", b"%PDF-1.4 garbage")
    assert extract_file(str(bad_pdf)).text_blocks == []


def test_dataless_file_is_not_read(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    p = touch(tmp_path / "cloud.txt", "내용")
    real = os.stat

    class Placeholder:
        st_flags = scan_mod.SF_DATALESS

    monkeypatch.setattr(
        extract_mod.os, "stat", lambda q, *a, **k: Placeholder() if str(q) == str(p) else real(q, *a, **k)
    )
    ex = extract_file(str(p))
    assert ex.extra["dataless"] is True and ex.text_blocks == []


def test_nfd_filename_title_is_nfc(tmp_path: Path) -> None:
    name = unicodedata.normalize("NFD", "한글파일") + ".txt"
    p = touch(tmp_path / name, "내용")
    assert extract_file(str(p)).title == "한글파일"


# ---------------------------------------------------------------- chunking


def _korean_doc(n_paras: int = 30, per_para: int = 5) -> str:
    paras = []
    k = 0
    for _ in range(n_paras):
        sents = []
        for _ in range(per_para):
            k += 1
            sents.append(f"{k}번째 문장은 한국어 청크 분할을 시험하기 위한 내용입니다.")
        paras.append(" ".join(sents))
    return "\n\n".join(paras)


def _overlap(a: str, b: str) -> int:
    for k in range(min(len(a), len(b)), 0, -1):
        if a.endswith(b[:k]):
            return k
    return 0


def test_chunk_sizes_overlap_and_sentence_boundaries() -> None:
    text = _korean_doc()
    chunks = chunk_blocks([TextBlock(text)])
    body = [c for c in chunks if c.field == "body"]
    assert len(body) > 3
    assert all(len(c.text) <= 900 for c in body)
    assert all(len(c.text) >= 400 for c in body[:-1])
    for a, b in zip(body, body[1:]):
        assert a.text.endswith("다.")  # cut at a sentence end
        ov = _overlap(a.text, b.text)
        assert 0 < ov <= 120
        assert b.text[:1].isdigit()  # overlap starts at a sentence start
    joined = " ".join(c.text for c in body)
    for k in (1, 77, 150):
        assert f"{k}번째 문장은" in joined


def test_chunk_long_paragraph_splits_on_korean_sentence_ends() -> None:
    sents = [f"문장{i}에서 회의 결과를 정리했다." if i % 2 else f"문장{i}도 확인해 주세요?" for i in range(80)]
    text = "".join(sents)  # no spaces after the full stops: Hangul + '.' still ends a sentence
    body = [c for c in chunk_blocks([TextBlock(text)], max_chars=300, overlap=60) if c.field == "body"]
    assert len(body) > 3 and all(len(c.text) <= 300 for c in body)
    for c in body[:-1]:
        assert c.text.endswith(("다.", "요?"))


def test_chunk_never_splits_words() -> None:
    vocab = ["개인", "통합", "인덱스", "검색", "파일", "문서", "tidy-one", "embedding", "서버", "맥북"]
    rng = random.Random(7)
    words = [rng.choice(vocab) for _ in range(1500)]
    body = [c for c in chunk_blocks([TextBlock(" ".join(words))]) if c.field == "body"]
    assert len(body) > 5
    for c in body:
        assert len(c.text) <= 900
        assert set(c.text.split()) <= set(vocab)
    for a, b in zip(body, body[1:]):
        assert 0 < _overlap(a.text, b.text) <= 120


def test_chunk_lines_and_hard_cut() -> None:
    rows = [f"{i},홍길동{i},서울시 강남구" for i in range(400)]
    body = [c for c in chunk_blocks([TextBlock("\n".join(rows))]) if c.field == "body"]
    row_set = set(rows)
    for c in body:
        assert len(c.text) <= 900
        assert all(line in row_set for line in c.text.split("\n"))
    blob = "가" * 2500
    pieces = split_text(blob, 780)
    assert "".join(pieces) == blob and max(map(len, pieces)) <= 780
    body = [c for c in chunk_blocks([TextBlock(blob)]) if c.field == "body"]
    assert all(len(c.text) <= 900 for c in body) and len(body) >= 3


def test_split_text_is_lossless() -> None:
    text = _korean_doc(8) + "\n\nEnglish sentence one. Second one? Third!\n" + "x" * 2000
    pieces = split_text(text, 300)
    assert "".join(pieces) == text and all(len(p) <= 300 for p in pieces)


def test_chunk_pages_and_meta(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("HOME", str(tmp_path))
    path = str(tmp_path / "Documents" / "회사" / "2026" / "분기 보고서.pdf")
    blocks = [
        TextBlock("첫 페이지 내용입니다.", 1),
        TextBlock(_korean_doc(10), 2),
        TextBlock("둘째 페이지 이어짐.", 2),
        TextBlock("셋째 페이지.", 3),
    ]
    modified = dt.datetime(2026, 10, 3, 14, 30)
    created = dt.datetime(2025, 3, 9, 9, 0).timestamp()
    chunks = chunk_blocks(blocks, title="분기 보고서", path=path, created=created, modified=modified,
                          doc_id="file:mba:abc")
    body = [c for c in chunks if c.field == "body"]
    meta = [c for c in chunks if c.field == "meta"]
    assert [c.ordinal for c in chunks] == list(range(len(chunks)))
    assert body[0].page == 1 and body[0].text == "첫 페이지 내용입니다."
    assert {c.page for c in body} == {1, 2, 3}
    assert [c.page for c in body] == sorted(c.page for c in body)
    assert "둘째 페이지 이어짐." in [c for c in body if c.page == 2][-1].text
    assert len(meta) == 1 and meta[0] is chunks[-1] and meta[0].page is None
    m = meta[0].text
    assert "제목: 분기 보고서" in m
    assert "폴더: Documents(문서) / 회사 / 2026" in m and "파일: 분기 보고서.pdf" in m
    assert "수정: 2026년 10월 3일 (2026-10-03)" in m and "생성: 2025년 3월 9일" in m
    assert str(tmp_path) not in m
    assert meta[0].chunk_id == f"file:mba:abc#meta:{len(body)}"
    assert body[0].chunk_id == "file:mba:abc#body:0"
    assert all(len(c.chunk_hash) == 32 for c in chunks)


def test_chunk_meta_only_for_media_and_icloud_path(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("HOME", str(tmp_path))
    path = str(tmp_path / "Library/Mobile Documents/com~apple~CloudDocs/여행/IMG_1.HEIC")
    chunks = chunk_blocks([], title="IMG_1", path=path, modified="2026-10-15T03:00:00Z")
    assert len(chunks) == 1 and chunks[0].field == "meta" and chunks[0].ordinal == 0
    assert "폴더: iCloud Drive / 여행" in chunks[0].text and "2026년 10월" in chunks[0].text
    assert chunks[0].chunk_id == "" and chunks[0].chunk_hash
    assert chunk_blocks([]) == []


def test_chunks_embed_with_fake_encoder(tmp_path: Path) -> None:
    p = touch(tmp_path / "doc.md", "# 예산 계획\n\n" + _korean_doc(6))
    ex = extract_file(str(p))
    chunks = chunk_blocks(ex.text_blocks, title=ex.title, path=str(p), doc_id="file:t:1")
    vecs = FakeEncoder().encode_passages([c.text for c in chunks])
    assert vecs.shape == (len(chunks), 1024)
    assert chunks[-1].field == "meta" and "예산 계획" in chunks[-1].text


# ---------------------------------------------------------------- regressions


def test_rtf_astral_characters_and_lone_surrogates(tmp_path: Path) -> None:
    """RTF writes U+1F600 as two \\uN surrogate halves; the text must hash and chunk, not fail every run."""

    def u(*codes: int) -> str:  # RTF unicode escapes: backslash, 'u', signed 16-bit decimal
        return "".join(f"{BS}u{c} " for c in codes)

    smile = u(-10179, -8704)  # U+1F600 as a UTF-16 surrogate pair
    rtf = (
        "{" + BS + "rtf1" + BS + "ansi" + BS + "ansicpg949 {" + BS + "info{" + BS + "title " + BS + "uc0"
        + u(54924) + smile + "}}" + BS + "uc0" + u(54924, 51032, 47197) + smile + " abc " + u(-10179) + "lone"
        + BS + "par}"
    )
    text, title = rtf_to_text(rtf)
    assert text.startswith("회의록😀 abc") and "\ufffdlone" in text and title == "회😀"
    ex = extract_file(str(touch(tmp_path / "memo.rtf", rtf)))
    assert "😀" in ex.text and ex.title == "회😀"
    chunks = chunk_blocks(ex.text_blocks, title=ex.title, doc_id="file:t:rtf")
    assert all(c.chunk_hash for c in chunks)
    assert extract_mod.clean_text("a\ud83d\ude00b\udc00c") == "a😀b\ufffdc"


def test_passport_pattern_is_linear_on_blank_runs() -> None:
    import time

    t = time.perf_counter()
    for label in ("여권", "Passport", "여권번호", "passport no"):
        assert find_passports(label + " " * 5000 + "x") == []
        assert find_passports((label + " " * 400 + "x ") * 20) == []
    assert time.perf_counter() - t < 1.0
    assert find_passports("여권 번호 :\nM12345678") == ["M12345678"]
    assert find_passports("Passport #M123A4567") == ["M123A4567"]


@pytest.mark.parametrize(
    "text",
    [
        "네이버 아이디: junwon\n비밀번호: Hunter2!@#",
        "password: hunter22",
        "PW = qwer1234",
        "와이파이 암호 : 12345678",
        "DATABASE_URL=postgres://app:s3cr3tPass@db.example.com:5432/app",
        "https://junwon:ghtoken123@github.com/org/repo.git",
        "GITLAB_TOKEN=glpat-AbCdEfGhIjKlMnOpQrSt",
        "npm_" + "a1B2c3D4e5" * 4,
        "SG." + "abcdEFGHijklMNOPqrstUV" + "." + "x" * 43,
        "eyJhbGciOiJIUzI1NiJ9.eyJzdWIiOiIxMjM0NTY3ODkwIn0.dozjgNryP4J3jVmNHl0w5N_XgL0n3I9PlFUP0THsR8U",
        "bot 123456789:" + "AAHdqTcvCH1vGWJxfSeofSAs0K5PALDsaw1",
    ],
)
def test_credential_lines_and_tokens_are_secret(text: str) -> None:
    assert content_sensitivity(text) & SENS_SECRET


@pytest.mark.parametrize(
    "text",
    [
        "Password: must be at least 8 characters",
        "비밀번호: 영문, 숫자 조합 8자 이상",
        "비밀번호 변경 안내",
        "Password: ********",
        "connect to postgres://user:password@localhost/db",
        "암호화 방식: AES-256",
        "회의는 10:30에 시작",
    ],
)
def test_credential_patterns_spare_ordinary_text(text: str) -> None:
    assert not content_sensitivity(text) & SENS_SECRET


def test_folders_named_like_secrets_are_not_walked(tmp_path: Path) -> None:
    root = tmp_path.resolve()
    for folder in ("암호", "tokens", "API Keys", "비밀번호", "Passwords", "Server Keys", "보통 폴더"):
        touch(root / folder / "bank.txt", "국민은행 계좌 정리")
    assert _rel(root, scan([str(root)])) == {"보통 폴더/bank.txt"}


def _tokens(texts: list[str]) -> list[int]:
    """Stand-in tokenizer: one token per Hangul syllable, one per 4 other characters, plus [CLS] and [SEP]."""
    return [2 + sum(1 for ch in t if "가" <= ch <= "힣") + sum(1 for ch in t if not "가" <= ch <= "힣") // 4
            for t in texts]


def test_chunks_are_split_to_fit_the_model_tokens() -> None:
    dense = "".join(f"근로계약서표준양식개정안검토결과보고서{i}번. " for i in range(200))
    english = "The quarterly review notes were shared with the whole team today. " * 60
    plain = [c.text for c in chunk_blocks([TextBlock(english)]) if c.field == "body"]
    fitted = [c for c in chunk_blocks([TextBlock(english)], count_tokens=_tokens, max_tokens=512) if c.field == "body"]
    assert [c.text for c in fitted] == plain  # chunks that fit are left alone (same hashes, cached vectors)
    before = [c.text for c in chunk_blocks([TextBlock(dense)]) if c.field == "body"]
    assert max(_tokens(before)) > 512
    after = chunk_blocks([TextBlock(dense)], count_tokens=_tokens, max_tokens=512, title="계약",
                         doc_id="file:t:dense")
    body = [c for c in after if c.field == "body"]
    assert max(_tokens([c.text for c in body])) <= 512 and len(body) > len(before)
    assert [c.ordinal for c in after] == list(range(len(after))) and after[-1].field == "meta"
    joined = " ".join(c.text for c in body)
    assert all(f"보고서{i}번." in joined for i in range(200))


def test_repeated_prefix_recovers_the_chunker_overlap_exactly() -> None:
    from tidy_one.extract.chunk import repeated_prefix

    dense = "".join(f"근로계약서표준양식개정안검토결과보고서{i}번. " for i in range(200))
    for text, kw in (
        ("출근,09:00,18:00,정상\n" * 200, {}),
        ("같은 문장이 반복된다. " * 150, {}),
        (_korean_doc(), {}),
        (dense, {"count_tokens": _tokens, "max_tokens": 512}),
    ):
        body = [c.text for c in chunk_blocks([TextBlock(text)], **kw) if c.field == "body"]
        merged = body[0]
        for nxt in body[1:]:
            n = repeated_prefix(merged, nxt)
            assert n is not None and n > 0
            merged += nxt[n:]
        assert merged == text.strip()
